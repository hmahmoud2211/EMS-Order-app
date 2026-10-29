"""Consolidated order exports: PDF, Excel, CSV and copyable text.

Exports describe an *internal* EMS order to be placed with the restaurant
externally; they never claim the restaurant accepted it. Fees are labelled as
admin-entered values.
"""

import csv
import io
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.clock import to_business_tz
from app.core.config import get_settings
from app.schemas.admin import ConsolidatedLine, ConsolidatedResponse
from app.schemas.order import SelectedModifierOut
from app.utils.text import contains_arabic

log = structlog.get_logger(__name__)

NAVY = colors.HexColor("#0B1736")
BLUE = colors.HexColor("#185ADB")
CYAN = colors.HexColor("#18B7D8")
INK = colors.HexColor("#17233B")
MUTED = colors.HexColor("#78869B")
LINE = colors.HexColor("#E4E9F2")
CANVAS = colors.HexColor("#F7F9FC")

LOGO_PATH = Path(__file__).resolve().parents[1] / "assets" / "logo.png"

FONT_CANDIDATES = [
    Path(__file__).resolve().parents[1] / "assets" / "fonts",
    Path("/usr/share/fonts/truetype/dejavu"),
]

_fonts: tuple[str, str] | None = None


def _register_fonts() -> tuple[str, str]:
    """Register a Unicode font with Arabic glyphs if available; fall back to Helvetica."""
    global _fonts
    if _fonts:
        return _fonts
    for folder in FONT_CANDIDATES:
        regular, bold = folder / "DejaVuSans.ttf", folder / "DejaVuSans-Bold.ttf"
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("EMSSans", str(regular)))
            pdfmetrics.registerFont(TTFont("EMSSans-Bold", str(bold)))
            _fonts = ("EMSSans", "EMSSans-Bold")
            return _fonts
    log.warning("pdf_unicode_font_missing", detail="Arabic text will not render in PDF exports")
    _fonts = ("Helvetica", "Helvetica-Bold")
    return _fonts


def _shape(text: str | None) -> str:
    """Shape and reorder Arabic so it renders correctly in the PDF."""
    if not text:
        return ""
    if contains_arabic(text):
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display

            return get_display(arabic_reshaper.reshape(text))
        except Exception:  # pragma: no cover - shaping libs are installed
            return text
    return text


def _p(text: str | None) -> str:
    return escape(_shape(text))


def money_text(value) -> str:
    return f"{value:,.2f}"


def modifiers_text(modifiers: list[SelectedModifierOut]) -> str:
    return ", ".join(f"{m.name} (+{m.price:,.2f})" if m.price and m.price > 0 else m.name for m in modifiers)


def _filename(data: ConsolidatedResponse, extension: str) -> str:
    safe = (
        "".join(c for c in data.summary.restaurant_name if c.isalnum() or c in " -_")
        .strip()
        .replace(" ", "-")
    )
    return f"EMS-Food-{data.summary.order_date.isoformat()}-{safe or 'order'}.{extension}"


def _local(dt: datetime) -> str:
    return to_business_tz(dt).strftime("%d %b %Y, %I:%M %p")


# ---- PDF -------------------------------------------------------------------------


def build_pdf(data: ConsolidatedResponse, company_name: str) -> tuple[bytes, str]:
    regular, bold = _register_fonts()
    currency = get_settings().currency
    s = data.summary
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=f"EMS Food order — {s.restaurant_name} — {s.order_date.isoformat()}",
        author=company_name,
    )
    base = ParagraphStyle("base", fontName=regular, fontSize=9, leading=12, textColor=INK)
    small = ParagraphStyle("small", parent=base, fontSize=8, leading=10, textColor=MUTED)
    h2 = ParagraphStyle(
        "h2", parent=base, fontName=bold, fontSize=12.5, leading=16, spaceBefore=8, spaceAfter=6
    )
    cell_b = ParagraphStyle("cellb", parent=base, fontName=bold)
    right = ParagraphStyle("right", parent=base, alignment=TA_RIGHT)
    right_b = ParagraphStyle("rightb", parent=right, fontName=bold)
    white_title = ParagraphStyle(
        "wt", parent=base, fontName=bold, fontSize=18, leading=22, textColor=colors.white
    )
    white_sub = ParagraphStyle(
        "ws", parent=base, fontSize=9.5, leading=13, textColor=colors.HexColor("#C9D6F2")
    )

    story = []
    brand = Paragraph("EMS <font color='#18B7D8'>Food</font>", white_title)
    if LOGO_PATH.exists():
        logo = Image(str(LOGO_PATH), width=11 * mm, height=11 * mm)
        brand_cell = Table([[logo, brand]], colWidths=[14 * mm, None])
        brand_cell.setStyle(
            TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)])
        )
    else:
        brand_cell = brand
    header = Table(
        [
            [
                brand_cell,
                Paragraph(
                    f"<b>Consolidated lunch order</b><br/>{escape(s.order_date.strftime('%A, %d %B %Y'))}",
                    ParagraphStyle("hr", parent=white_sub, alignment=TA_RIGHT, textColor=colors.white),
                ),
            ],
            [Paragraph(escape(company_name), white_sub), Paragraph("", white_sub)],
        ],
        colWidths=[doc.width * 0.55, doc.width * 0.45],
    )
    header.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), NAVY),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, 0), 12),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 12),
                ("LINEBELOW", (0, -1), (-1, -1), 3, CYAN),
            ]
        )
    )
    story += [header, Spacer(1, 10)]

    restaurant = s.restaurant_name + (f" — {s.restaurant_branch}" if s.restaurant_branch else "")
    summary_rows = [
        ["Restaurant", _p(restaurant), "Employees ordered", str(s.employees_ordered)],
        ["Restaurant phone", _p(s.restaurant_phone or "—"), "Meals", str(s.meals)],
        [
            "Session status",
            s.session_status.title(),
            "Food subtotal",
            f"{money_text(s.food_subtotal)} {currency}",
        ],
    ]
    summary = Table(
        [
            [Paragraph(str(c), cell_b if i % 2 == 0 else base) for i, c in enumerate(row)]
            for row in summary_rows
        ],
        colWidths=[doc.width * 0.18, doc.width * 0.37, doc.width * 0.2, doc.width * 0.25],
    )
    summary.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), CANVAS),
                ("BOX", (0, 0), (-1, -1), 0.6, LINE),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, LINE),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story += [summary, Spacer(1, 6)]

    # Totals block
    totals = [
        [Paragraph("Food subtotal", base), Paragraph(f"{money_text(s.food_subtotal)} {currency}", right)]
    ]
    if s.delivery_fee > 0:
        totals.append(
            [
                Paragraph("Delivery fee (entered by admin)", base),
                Paragraph(f"{money_text(s.delivery_fee)} {currency}", right),
            ]
        )
    if s.service_fee > 0:
        totals.append(
            [
                Paragraph("Service fee (entered by admin)", base),
                Paragraph(f"{money_text(s.service_fee)} {currency}", right),
            ]
        )
    totals.append(
        [
            Paragraph("<b>Grand total</b>", cell_b),
            Paragraph(f"<b>{money_text(s.final_total)} {currency}</b>", right_b),
        ]
    )
    totals_table = Table(totals, colWidths=[doc.width * 0.7, doc.width * 0.3])
    totals_table.setStyle(
        TableStyle(
            [
                ("LINEABOVE", (0, -1), (-1, -1), 1, BLUE),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story += [totals_table]
    if s.fee_note:
        story.append(Paragraph(f"Fee note (admin): {_p(s.fee_note)}", small))

    # By item — what to tell the restaurant
    story.append(Paragraph("Order for the restaurant (grouped by item)", h2))
    item_rows = [[Paragraph(f"<b>{h}</b>", base) for h in ("Qty", "Item & options", "For", "Unit", "Total")]]
    for item in data.by_item:
        options = modifiers_text(item.modifiers)
        who = "<br/>".join(
            f"{_p(e.full_name)} ×{e.quantity}"
            + (
                f" <font color='#B45309'>— {_p(e.special_instructions)}</font>"
                if e.special_instructions
                else ""
            )
            for e in item.employees
        )
        item_rows.append(
            [
                Paragraph(f"<b>{item.total_quantity}</b>", base),
                Paragraph(
                    f"<b>{_p(item.item_name)}</b>"
                    + (f"<br/><font size=8 color='#3A4760'>{_p(options)}</font>" if options else ""),
                    base,
                ),
                Paragraph(who, small),
                Paragraph(money_text(item.unit_price), right),
                Paragraph(money_text(item.total_value), right),
            ]
        )
    if len(item_rows) == 1:
        item_rows.append([Paragraph("No orders were submitted.", small), "", "", "", ""])
    items_table = Table(
        item_rows,
        colWidths=[doc.width * 0.07, doc.width * 0.38, doc.width * 0.31, doc.width * 0.11, doc.width * 0.13],
        repeatRows=1,
    )
    items_table.setStyle(_grid_style())
    story.append(items_table)

    # By employee
    story.append(Paragraph("Orders by employee", h2))
    for emp in data.by_employee:
        rows = [
            [
                Paragraph(
                    f"<b>{_p(emp.full_name)}</b>"
                    + (f" <font color='#78869B'>· {_p(emp.department)}</font>" if emp.department else "")
                    + f" <font color='#78869B'>· {escape(emp.order_reference)}</font>",
                    base,
                ),
                Paragraph(f"<b>{money_text(emp.total)} {currency}</b>", right_b),
            ]
        ]
        for line in emp.lines:
            rows.append([Paragraph(_line_markup(line), base), Paragraph(money_text(line.line_total), right)])
        if emp.allocated_fees > 0:
            rows.append(
                [
                    Paragraph("Share of admin-entered fees", small),
                    Paragraph(money_text(emp.allocated_fees), right),
                ]
            )
        block = Table(rows, colWidths=[doc.width * 0.78, doc.width * 0.22])
        block.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), CANVAS),
                    ("BOX", (0, 0), (-1, -1), 0.6, LINE),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.6, LINE),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story += [KeepTogether([block, Spacer(1, 5)])]

    generated = _local(data.generated_at)

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(
            doc_.leftMargin,
            9 * mm,
            f"Generated {generated} (Cairo) · "
            "Internal EMS order summary — not a restaurant or Talabat confirmation",
        )
        canvas.drawRightString(A4[0] - doc_.rightMargin, 9 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue(), _filename(data, "pdf")


def _line_markup(line: ConsolidatedLine) -> str:
    text = f"{line.quantity} × <b>{_p(line.item_name)}</b>"
    options = modifiers_text(line.modifiers)
    if options:
        text += f" <font color='#3A4760'>({_p(options)})</font>"
    if line.special_instructions:
        text += f"<br/><font color='#B45309'>Note: {_p(line.special_instructions)}</font>"
    return text


def _grid_style() -> TableStyle:
    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), CANVAS),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, BLUE),
            ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
    )


# ---- Excel -----------------------------------------------------------------------


def build_xlsx(data: ConsolidatedResponse, company_name: str) -> tuple[bytes, str]:
    s = data.summary
    currency = get_settings().currency
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0B1736")
    thin = Side(style="thin", color="E4E9F2")
    money_fmt = "#,##0.00"

    def style_header(ws, row: int, count: int) -> None:
        for col in range(1, count + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "EMS Food — Consolidated lunch order"
    ws["A1"].font = Font(bold=True, size=14, color="0B1736")
    ws["A2"] = company_name
    ws["A2"].font = Font(color="78869B")
    rows = [
        ("Date", s.order_date.isoformat()),
        ("Restaurant", s.restaurant_name),
        ("Branch", s.restaurant_branch or ""),
        ("Restaurant phone", s.restaurant_phone or ""),
        ("Session status", s.session_status),
        ("Employees ordered", s.employees_ordered),
        ("Meals", s.meals),
        (f"Food subtotal ({currency})", float(s.food_subtotal)),
        (f"Delivery fee — admin-entered ({currency})", float(s.delivery_fee)),
        (f"Service fee — admin-entered ({currency})", float(s.service_fee)),
        ("Fee allocation", s.fee_allocation),
        ("Fee note", s.fee_note or ""),
        (f"Grand total ({currency})", float(s.final_total)),
        ("Generated (Cairo)", _local(data.generated_at)),
        ("Note", "Internal EMS order summary — not a restaurant confirmation."),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=value)
        if isinstance(value, float):
            cell.number_format = money_fmt
    ws.cell(row=4 + len(rows) - 3, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 44

    ws2 = wb.create_sheet("By Item")
    headers = [
        "Quantity",
        "Item",
        "Category",
        "Options",
        f"Unit price ({currency})",
        f"Total ({currency})",
        "Orders",
        "Employees",
        "Special instructions",
    ]
    ws2.append(headers)
    style_header(ws2, 1, len(headers))
    for item in data.by_item:
        ws2.append(
            [
                item.total_quantity,
                item.item_name,
                item.category_name or "",
                modifiers_text(item.modifiers),
                float(item.unit_price),
                float(item.total_value),
                item.order_count,
                ", ".join(f"{e.full_name} ×{e.quantity}" for e in item.employees),
                "; ".join(
                    f"{e.full_name}: {e.special_instructions}"
                    for e in item.employees
                    if e.special_instructions
                ),
            ]
        )
    _finish_sheet(ws2, [10, 32, 18, 36, 14, 14, 9, 40, 40], money_cols=[5, 6], money_fmt=money_fmt)

    ws3 = wb.create_sheet("By Employee")
    headers = [
        "Employee",
        "Department",
        "Order ref",
        "Quantity",
        "Item",
        "Options",
        "Special instructions",
        f"Unit price ({currency})",
        f"Line total ({currency})",
        f"Employee subtotal ({currency})",
        f"Fee share ({currency})",
        f"Employee total ({currency})",
    ]
    ws3.append(headers)
    style_header(ws3, 1, len(headers))
    for emp in data.by_employee:
        for idx, line in enumerate(emp.lines):
            first = idx == 0
            ws3.append(
                [
                    emp.full_name,
                    emp.department or "",
                    emp.order_reference,
                    line.quantity,
                    line.item_name,
                    modifiers_text(line.modifiers),
                    line.special_instructions or "",
                    float(line.unit_price),
                    float(line.line_total),
                    float(emp.subtotal) if first else None,
                    float(emp.allocated_fees) if first else None,
                    float(emp.total) if first else None,
                ]
            )
    total_row = ws3.max_row + 1
    ws3.cell(row=total_row, column=1, value="Grand total").font = Font(bold=True)
    grand = ws3.cell(row=total_row, column=12, value=float(s.final_total))
    grand.font = Font(bold=True)
    _finish_sheet(
        ws3,
        [24, 16, 15, 9, 30, 34, 30, 12, 12, 14, 12, 14],
        money_cols=[8, 9, 10, 11, 12],
        money_fmt=money_fmt,
    )
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        for cell in row:
            cell.border = Border(bottom=thin)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), _filename(data, "xlsx")


def _finish_sheet(ws, widths: list[int], money_cols: list[int], money_fmt: str) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = money_fmt
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---- CSV / text ---------------------------------------------------------------------


def build_csv(data: ConsolidatedResponse) -> tuple[bytes, str]:
    s = data.summary
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(
        [
            "date",
            "restaurant",
            "employee",
            "department",
            "order_reference",
            "quantity",
            "item",
            "options",
            "special_instructions",
            "unit_price",
            "line_total",
            "employee_subtotal",
            "fee_share",
            "employee_total",
        ]
    )
    for emp in data.by_employee:
        for line in emp.lines:
            writer.writerow(
                [
                    s.order_date.isoformat(),
                    s.restaurant_name,
                    emp.full_name,
                    emp.department or "",
                    emp.order_reference,
                    line.quantity,
                    line.item_name,
                    modifiers_text(line.modifiers),
                    line.special_instructions or "",
                    f"{line.unit_price:.2f}",
                    f"{line.line_total:.2f}",
                    f"{emp.subtotal:.2f}",
                    f"{emp.allocated_fees:.2f}",
                    f"{emp.total:.2f}",
                ]
            )
    writer.writerow([])
    writer.writerow(["food_subtotal", f"{s.food_subtotal:.2f}"])
    writer.writerow(["delivery_fee_admin_entered", f"{s.delivery_fee:.2f}"])
    writer.writerow(["service_fee_admin_entered", f"{s.service_fee:.2f}"])
    writer.writerow(["grand_total", f"{s.final_total:.2f}"])
    # UTF-8 BOM so Excel opens Arabic names correctly.
    return ("﻿" + out.getvalue()).encode("utf-8"), _filename(data, "csv")


def build_text(data: ConsolidatedResponse, company_name: str) -> str:
    s = data.summary
    currency = get_settings().currency
    restaurant = s.restaurant_name + (f" ({s.restaurant_branch})" if s.restaurant_branch else "")
    lines = [
        f"EMS Food — {company_name}",
        f"Lunch order: {restaurant}",
        f"Date: {s.order_date.strftime('%A %d %B %Y')}",
        f"Employees: {s.employees_ordered} · Meals: {s.meals}",
        "",
        "ORDER FOR THE RESTAURANT",
    ]
    for item in data.by_item:
        options = modifiers_text(item.modifiers)
        lines.append(
            f"• {item.total_quantity} × {item.item_name}"
            + (f" ({options})" if options else "")
            + f" — {money_text(item.total_value)} {currency}"
        )
        notes = [f"{e.full_name}: {e.special_instructions}" for e in item.employees if e.special_instructions]
        for note in notes:
            lines.append(f"    note — {note}")
    lines += ["", "BY EMPLOYEE"]
    for emp in data.by_employee:
        lines.append(f"{emp.full_name} — {money_text(emp.total)} {currency}")
        for line in emp.lines:
            options = modifiers_text(line.modifiers)
            lines.append(
                f"   {line.quantity} × {line.item_name}"
                + (f" ({options})" if options else "")
                + (f" — note: {line.special_instructions}" if line.special_instructions else "")
            )
    lines += ["", f"Food subtotal: {money_text(s.food_subtotal)} {currency}"]
    if s.delivery_fee > 0:
        lines.append(f"Delivery fee (entered by admin): {money_text(s.delivery_fee)} {currency}")
    if s.service_fee > 0:
        lines.append(f"Service fee (entered by admin): {money_text(s.service_fee)} {currency}")
    lines.append(f"TOTAL: {money_text(s.final_total)} {currency}")
    return "\n".join(lines)
