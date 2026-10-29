import csv
import io
from decimal import Decimal

import pymupdf
from openpyxl import load_workbook

from app.services.report_service import _allocate_fees
from tests.conftest import login, make_user, open_session

API = "/api/v1"


def place(client, headers, sid, lines):
    response = client.post(f"{API}/orders", json={"session_id": sid, "items": lines}, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()


def scenario(client, db, admin_headers, menu):
    for name, full in [("ahmed", "Ahmed Kamal"), ("sara", "Sara Nabil"), ("omar", "Omar Fathy")]:
        make_user(db, name, full_name=full)
    sid = open_session(client, admin_headers, menu["restaurant"].id)["id"]
    burger_large_cheese = {
        "menu_item_id": str(menu["burger"].id),
        "quantity": 1,
        "modifier_ids": [str(menu["large"].id), str(menu["cheese"].id)],
        "special_instructions": "No onions",
    }
    burger_regular = {
        "menu_item_id": str(menu["burger"].id),
        "quantity": 1,
        "modifier_ids": [str(menu["regular"].id)],
    }
    place(
        client,
        login(client, "ahmed"),
        sid,
        [burger_large_cheese, {"menu_item_id": str(menu["fries"].id), "quantity": 1}],
    )
    place(
        client,
        login(client, "sara"),
        sid,
        [
            {"menu_item_id": str(menu["pasta"].id), "quantity": 1},
            {**burger_large_cheese, "special_instructions": None},
        ],
    )
    place(client, login(client, "omar"), sid, [burger_regular])
    return sid


def test_dashboard_and_employee_status(client, db, admin_headers, menu):
    make_user(db, "idle", full_name="Idle Person")
    sid = scenario(client, db, admin_headers, menu)
    dashboard = client.get(f"{API}/admin/dashboard", headers=admin_headers).json()
    metrics = dashboard["metrics"]
    assert metrics["total_employees"] == 4
    assert metrics["ordered"] == 3
    assert metrics["not_ordered"] == 1
    # Ahmed 235+60=295, Sara 185+235=420, Omar 180
    assert metrics["total_value"] == "895.00"
    assert metrics["total_items"] == 5
    assert dashboard["not_ordered"][0]["full_name"] == "Idle Person"
    assert dashboard["seconds_until_close"] == 7200

    status = client.get(
        f"{API}/admin/sessions/{sid}/employee-status", params={"filter": "not_ordered"}, headers=admin_headers
    ).json()
    assert status["counts"] == {"all": 4, "ordered": 3, "not_ordered": 1}
    assert [r["full_name"] for r in status["rows"]] == ["Idle Person"]
    searched = client.get(
        f"{API}/admin/sessions/{sid}/employee-status", params={"q": "sar"}, headers=admin_headers
    ).json()
    assert [r["full_name"] for r in searched["rows"]] == ["Sara Nabil"]


def test_consolidated_groups_identical_configurations(client, db, admin_headers, menu):
    sid = scenario(client, db, admin_headers, menu)
    client.patch(
        f"{API}/sessions/{sid}",
        json={"delivery_fee": "30.00", "fee_allocation": "EQUAL"},
        headers=admin_headers,
    )
    data = client.get(f"{API}/admin/sessions/{sid}/consolidated", headers=admin_headers).json()
    by_item = {(i["item_name"], tuple(m["name"] for m in i["modifiers"])): i for i in data["by_item"]}
    large_cheese = by_item[("Chicken Burger", ("Large", "Cheese"))]
    assert large_cheese["total_quantity"] == 2
    assert large_cheese["order_count"] == 2
    assert large_cheese["total_value"] == "470.00"
    assert {e["full_name"]: e["special_instructions"] for e in large_cheese["employees"]} == {
        "Ahmed Kamal": "No onions",
        "Sara Nabil": None,
    }
    assert by_item[("Chicken Burger", ("Regular",))]["total_quantity"] == 1

    summary = data["summary"]
    assert summary["employees_ordered"] == 3
    assert summary["meals"] == 5
    assert summary["food_subtotal"] == "895.00"
    assert summary["fees_total"] == "30.00"
    assert summary["fees_are_admin_entered"] is True
    assert summary["final_total"] == "925.00"
    assert [e["allocated_fees"] for e in data["by_employee"]] == ["10.00", "10.00", "10.00"]
    assert sum(Decimal(e["total"]) for e in data["by_employee"]) == Decimal("925.00")


def test_fee_allocation_rounding():
    assert _allocate_fees([Decimal("100"), Decimal("100"), Decimal("100")], Decimal("10.00"), "EQUAL") == [
        Decimal("3.34"),
        Decimal("3.33"),
        Decimal("3.33"),
    ]
    shares = _allocate_fees([Decimal("300"), Decimal("100")], Decimal("10.00"), "PROPORTIONAL")
    assert shares == [Decimal("7.50"), Decimal("2.50")]
    assert sum(
        _allocate_fees([Decimal("1"), Decimal("2"), Decimal("4")], Decimal("1.00"), "PROPORTIONAL")
    ) == Decimal("1.00")
    assert _allocate_fees([Decimal("100")], Decimal("10"), "NONE") == [Decimal("0.00")]


def test_exports_contain_names_items_and_totals(client, db, admin_headers, menu):
    sid = scenario(client, db, admin_headers, menu)
    pdf = client.get(f"{API}/admin/sessions/{sid}/export", params={"format": "pdf"}, headers=admin_headers)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert "attachment" in pdf.headers["content-disposition"]
    text = "".join(page.get_text() for page in pymupdf.open(stream=pdf.content, filetype="pdf"))
    for expected in (
        "Ahmed Kamal",
        "Sara Nabil",
        "Omar Fathy",
        "Chicken Burger",
        "Large",
        "No onions",
        "895.00",
        "Restaurant A",
    ):
        assert expected in text, expected
    assert "not a restaurant or Talabat confirmation" in text

    xlsx = client.get(f"{API}/admin/sessions/{sid}/export", params={"format": "xlsx"}, headers=admin_headers)
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Summary", "By Item", "By Employee"]
    by_employee = workbook["By Employee"]
    names = {row[0] for row in by_employee.iter_rows(min_row=2, values_only=True) if row[0]}
    assert {"Ahmed Kamal", "Sara Nabil", "Omar Fathy", "Grand total"} <= names
    assert by_employee.cell(row=by_employee.max_row, column=12).value == 895.0

    csv_resp = client.get(
        f"{API}/admin/sessions/{sid}/export", params={"format": "csv"}, headers=admin_headers
    )
    rows = list(csv.reader(io.StringIO(csv_resp.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["date", "restaurant", "employee"]
    assert ["grand_total", "895.00"] in rows

    txt = client.get(
        f"{API}/admin/sessions/{sid}/export", params={"format": "txt"}, headers=admin_headers
    ).text
    assert "2 × Chicken Burger (Large (+35.00), Cheese (+20.00))" in txt
    assert "TOTAL: 895.00 EGP" in txt


def test_analytics(client, db, admin_headers, menu):
    scenario(client, db, admin_headers, menu)
    spending = client.get(f"{API}/admin/analytics/spending", headers=admin_headers).json()
    assert spending["total_value"] == "895.00"
    assert spending["by_employee"][0]["full_name"] == "Sara Nabil"
    participation = client.get(f"{API}/admin/analytics/participation", headers=admin_headers).json()
    assert participation["by_day"][0]["ordered"] == 3
    assert participation["by_day"][0]["rate_pct"] == 100.0
