"""Spec §24 — final acceptance scenario, executed end to end over the HTTP API."""

import io
from datetime import time, timedelta

import pymupdf
from openpyxl import load_workbook

from app.core.clock import business_datetime
from app.utils import sample_menus
from tests.conftest import TEST_NOW, login

API = "/api/v1"


def change_password(client, username, temporary):
    headers = login(client, username, temporary)
    response = client.post(
        f"{API}/auth/change-password",
        json={"current_password": temporary, "new_password": "Lunch2026ok"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_full_acceptance_scenario(client, admin_headers, clock):
    # 1. Admin creates 18 employee accounts.
    temps = {}
    for i in range(1, 19):
        response = client.post(
            f"{API}/employees",
            json={
                "full_name": f"Employee {i:02d}",
                "username": f"employee{i:02d}",
                "department": "Engineering",
            },
            headers=admin_headers,
        )
        assert response.status_code == 201, response.text
        temps[i] = response.json()["temporary_password"]

    # 2. Admin creates a restaurant and imports its menu.
    rid = client.post(
        f"{API}/restaurants", json={"name": "Restaurant A", "branch": "Main"}, headers=admin_headers
    ).json()["id"]
    import_id = client.post(
        f"{API}/restaurants/{rid}/menu/import",
        files={"file": ("menu.csv", sample_menus.sample_csv(), "text/csv")},
        headers=admin_headers,
    ).json()["id"]
    approved = client.post(
        f"{API}/restaurants/{rid}/menu/imports/{import_id}/approve",
        json={"mode": "replace"},
        headers=admin_headers,
    )
    assert approved.json()["status"] == "APPROVED"

    # A second restaurant whose items must never be accepted today.
    other_id = client.post(f"{API}/restaurants", json={"name": "Restaurant B"}, headers=admin_headers).json()[
        "id"
    ]
    other_import = client.post(
        f"{API}/restaurants/{other_id}/menu/import",
        files={"file": ("menu.xlsx", sample_menus.sample_xlsx(), "application/octet-stream")},
        headers=admin_headers,
    ).json()["id"]
    client.post(
        f"{API}/restaurants/{other_id}/menu/imports/{other_import}/approve",
        json={"mode": "replace"},
        headers=admin_headers,
    )

    # 3–4. Admin selects that restaurant for today and opens ordering until 12:00.
    noon = business_datetime(TEST_NOW.date(), time(12, 0))
    session = client.post(
        f"{API}/sessions",
        json={"restaurant_id": rid, "closes_at": noon.isoformat(), "open_now": True},
        headers=admin_headers,
    ).json()
    sid = session["id"]
    assert session["status"] == "OPEN"

    e1 = change_password(client, "employee01", temps[1])
    e2 = change_password(client, "employee02", temps[2])
    e3 = change_password(client, "employee03", temps[3])

    # Employees see only today's restaurant and its menu.
    today = client.get(f"{API}/sessions/today", headers=e1).json()
    assert today["session"]["restaurant"]["name"] == "Restaurant A"
    assert today["is_accepting_orders"] is True
    menu = client.get(f"{API}/sessions/today/menu", headers=e1).json()
    items = {i["name"]: i for c in menu["categories"] for i in c["items"]}
    burger, pasta = items["Chicken Burger"], items["Chicken Pasta"]
    regular = next(m for m in burger["modifier_groups"][0]["modifiers"] if m["name"] == "Regular")

    # 5. Employee 1 orders a burger.
    order1 = client.post(
        f"{API}/orders",
        json={
            "session_id": sid,
            "items": [{"menu_item_id": burger["id"], "quantity": 1, "modifier_ids": [regular["id"]]}],
            "expected_total": "180.00",
        },
        headers={**e1, "Idempotency-Key": "e1-first-submit"},
    )
    assert order1.status_code == 201, order1.text
    # 6. Employee 2 orders pasta.
    order2 = client.post(
        f"{API}/orders",
        json={"session_id": sid, "items": [{"menu_item_id": pasta["id"], "quantity": 1}]},
        headers=e2,
    )
    assert order2.status_code == 201
    # 7. Employee 3 logs in and does not order.
    assert client.get(f"{API}/sessions/today", headers=e3).json()["my_order"] is None

    # 8. Dashboard: 18 active employees, 2 submitted, 16 without, correct total.
    metrics = client.get(f"{API}/admin/dashboard", headers=admin_headers).json()["metrics"]
    assert (metrics["total_employees"], metrics["ordered"], metrics["not_ordered"]) == (18, 2, 16)
    assert metrics["total_value"] == "365.00"  # 180 + 185

    # 9–10. Employee 1 edits the burger quantity before the deadline; the dashboard updates.
    edited = client.put(
        f"{API}/orders/{order1.json()['id']}",
        json={"items": [{"menu_item_id": burger["id"], "quantity": 2, "modifier_ids": [regular["id"]]}]},
        headers=e1,
    )
    assert edited.json()["total_amount"] == "360.00"
    metrics = client.get(f"{API}/admin/dashboard", headers=admin_headers).json()["metrics"]
    assert metrics["total_value"] == "545.00"
    assert metrics["ordered"] == 2

    # 11. An item from a different restaurant is rejected by the backend.
    other_menu = client.get(f"{API}/restaurants/{other_id}/menu", headers=admin_headers).json()
    foreign_item = other_menu["categories"][0]["items"][0]["id"]
    foreign = client.post(
        f"{API}/orders",
        json={"session_id": sid, "items": [{"menu_item_id": foreign_item, "quantity": 1}]},
        headers=e3,
    )
    assert foreign.status_code == 422
    assert foreign.json()["error"]["code"] == "ITEM_NOT_IN_MENU"

    # 12–13. Admin closes the session; a later submission is rejected.
    assert client.post(f"{API}/sessions/{sid}/close", headers=admin_headers).json()["status"] == "CLOSED"
    late = client.post(
        f"{API}/orders",
        json={"session_id": sid, "items": [{"menu_item_id": pasta["id"], "quantity": 1}]},
        headers=e3,
    )
    assert late.status_code == 409
    assert late.json()["error"]["code"] == "SESSION_CLOSED"

    # 14–15. Consolidated PDF and Excel contain names, configurations, quantities, prices and total.
    pdf = client.get(f"{API}/admin/sessions/{sid}/export", params={"format": "pdf"}, headers=admin_headers)
    text = "".join(page.get_text() for page in pymupdf.open(stream=pdf.content, filetype="pdf"))
    for expected in (
        "Employee 01",
        "Employee 02",
        "Chicken Burger",
        "Regular",
        "Chicken Pasta",
        "360.00",
        "185.00",
        "545.00",
    ):
        assert expected in text, expected
    assert "Employee 03" not in text

    workbook = load_workbook(
        io.BytesIO(
            client.get(
                f"{API}/admin/sessions/{sid}/export", params={"format": "xlsx"}, headers=admin_headers
            ).content
        )
    )
    by_item = {row[1]: row for row in workbook["By Item"].iter_rows(min_row=2, values_only=True)}
    assert by_item["Chicken Burger"][0] == 2  # quantity
    assert by_item["Chicken Burger"][3] == "Regular"  # configuration
    assert by_item["Chicken Burger"][5] == 360.0
    assert by_item["Chicken Pasta"][5] == 185.0
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=4, values_only=True)}
    assert summary["Grand total (EGP)"] == 545.0
    assert summary["Employees ordered"] == 2

    # 16. Admin finalizes; history remains available.
    assert (
        client.post(f"{API}/sessions/{sid}/finalize", headers=admin_headers).json()["status"] == "FINALIZED"
    )
    history = client.get(f"{API}/orders/me/history", headers=e1).json()
    assert history["items"][0]["total_amount"] == "360.00"
    assert history["items"][0]["restaurant_name"] == "Restaurant A"

    # 17. Next day uses a different restaurant without affecting the previous session.
    clock.set(TEST_NOW + timedelta(days=1))
    next_noon = business_datetime((TEST_NOW + timedelta(days=1)).date(), time(12, 0))
    next_session = client.post(
        f"{API}/sessions",
        json={"restaurant_id": other_id, "closes_at": next_noon.isoformat(), "open_now": True},
        headers=admin_headers,
    )
    assert next_session.status_code == 201, next_session.text
    new_sid = next_session.json()["id"]
    assert (
        client.get(f"{API}/sessions/today", headers=e1).json()["session"]["restaurant"]["name"]
        == "Restaurant B"
    )
    assert client.get(f"{API}/orders/me/today", headers=e1).json() is None
    # Yesterday's restaurant items are rejected today.
    stale = client.post(
        f"{API}/orders",
        json={"session_id": new_sid, "items": [{"menu_item_id": burger["id"], "quantity": 1}]},
        headers=e1,
    )
    assert stale.json()["error"]["code"] == "ITEM_NOT_IN_MENU"
    # The previous session is untouched.
    previous = client.get(f"{API}/admin/sessions/{sid}/consolidated", headers=admin_headers).json()
    assert previous["summary"]["final_total"] == "545.00"
    assert previous["session"]["status"] == "FINALIZED"
    assert client.get(f"{API}/orders/me/history", headers=e1).json()["total"] == 1
